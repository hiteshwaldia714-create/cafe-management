import sqlite3
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from datetime import datetime


def export_to_excel():

    # Connect to database
    connection = sqlite3.connect("cafe.db")
    cursor = connection.cursor()

    # Get all completed sessions
    cursor.execute("""
        SELECT
            sessions.customer_id,
            customers.name,
            sessions.pc_name,
            sessions.login_time,
            sessions.logout_time
        FROM sessions
        JOIN customers
        ON sessions.customer_id = customers.customer_id
        WHERE sessions.logout_time IS NOT NULL
        ORDER BY sessions.id
    """)

    sessions = cursor.fetchall()

    connection.close()


    # Create Excel workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cafe Sessions"


    # Create headers
    headers = [
        "Customer ID",
        "Name",
        "PC",
        "Login Time",
        "Logout Time",
        "Duration"
    ]

    sheet.append(headers)


    # Make headers bold
    for cell in sheet[1]:
        cell.font = Font(bold=True)


    # Add session data
    for session in sessions:

        customer_id = session[0]
        name = session[1]
        pc_name = session[2]
        login_time = session[3]
        logout_time = session[4]

        login = datetime.strptime(
            login_time,
            "%Y-%m-%d %H:%M"
        )

        logout = datetime.strptime(
            logout_time,
            "%Y-%m-%d %H:%M"
        )

        duration = logout - login

        sheet.append([
            customer_id,
            name,
            pc_name,
            login_time,
            logout_time,
            str(duration)
        ])


    # Adjust column widths
    for column in sheet.columns:

        max_length = 0
        column_letter = get_column_letter(column[0].column)

        for cell in column:

            if cell.value:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        sheet.column_dimensions[column_letter].width = max_length + 2


    # Save Excel file
    workbook.save("Cafe_Sessions.xlsx")

    print("Excel file updated successfully!")

if __name__ == "__main__":
    export_to_excel()